from io import BytesIO
from openpyxl import Workbook
from database import connect
from photo_match import match_column
from desktop_data import result_filter


def export_xlsx(language='en',media_type='all',match_kind='auto'):
    column=match_column(match_kind,media_type)
    if media_type not in ('all','photo','video'):
        raise ValueError('Invalid media type')
    chinese=language=='zh'
    wb=Workbook(write_only=True)
    ws=wb.create_sheet(('视频完全重复' if chinese else 'Video duplicates') if media_type=='video' else ('完全重复' if chinese else 'Exact duplicates'))
    if match_kind=='pixels':ws.title='相同照片（忽略元数据）' if chinese else 'Same photo - metadata ignored'
    if match_kind=='auto':ws.title='重复媒体' if chinese else 'Duplicate media'
    ws.append(['组 ID','副本序号','完整路径','文件大小（字节）','根文件夹','多余副本（仅供参考）'] if chinese else ['Group ID','Copy #','Full Path','File Size (bytes)','Root Folder','Potential Extra Copy'])
    c=connect()
    condition,parameters=result_filter(media_type)
    rows=c.execute(f'SELECT p.{column} sha256,p.path,p.size,p.root FROM photos p JOIN (SELECT {column} digest FROM photos WHERE active=1 AND {column} IS NOT NULL AND '+condition+f' GROUP BY {column} HAVING count(*)>1) d ON p.{column}=d.digest WHERE p.active=1 AND '+condition.replace('media_type','p.media_type')+f' ORDER BY p.{column},p.path',parameters+parameters)
    old=None
    i=0
    for r in rows:
        if r['sha256']!=old: old=r['sha256']; i=0
        i+=1
        ws.append([r['sha256'],i,r['path'],r['size'],r['root'],('否' if i==1 else '是（仅供检查）') if chinese else ('No' if i==1 else 'Yes (review only)')])
    if match_kind=='auto' and media_type in ('photo','all'):
        similar=wb.create_sheet('Similar photos' if not chinese else '\u76f8\u4f3c\u7167\u7247')
        similar.append(['Group ID','Full Path','File Size (bytes)','Visual score (%)','Review required'])
        for r in c.execute('SELECT s.group_id,p.path,p.size,s.score FROM similarity_members s JOIN photos p ON p.path=s.path WHERE p.active=1 AND s.group_id IN (SELECT s.group_id FROM similarity_members s JOIN photos p ON p.path=s.path WHERE p.active=1 GROUP BY s.group_id HAVING count(*)>1) ORDER BY s.group_id,p.path'):
            similar.append([r['group_id'],r['path'],r['size'],r['score'],True])
    err=wb.create_sheet('扫描错误' if chinese else 'Scan errors')
    err.append(['扫描编号','路径','错误'] if chinese else ['Run ID','Path','Error'])
    for r in c.execute('SELECT run_id,path,error FROM scan_errors ORDER BY id DESC LIMIT 20000'):
        err.append([r['run_id'],r['path'],r['error']])
    c.close()
    out=BytesIO()
    wb.save(out)
    return out.getvalue()
